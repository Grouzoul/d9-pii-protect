#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
Moteur d'anonymisation gouvernée des PII (sidecar) — Presidio FR embarqué.

  détection = Presidio (analyzer.py) : NER FR Anonym-IA + recognizers regex/checksum + scoring contextuel
  chiffrement réversible = AES-256-GCM in-place, marqueurs [[PII:TYPE:b64]] (crypto.py)

Endpoints (JSON) :
  POST /anonymize      {text}                  -> {text, count, entities}   (chiffre les PII détectées)
  POST /encrypt        {text}                  -> {text, count}             (chiffre la valeur ENTIÈRE, 1 marqueur)
  POST /rehydrate      {text}                  -> {text, count}
  POST /anonymize-doc  {filename, content_b64} -> {filename, content_b64}   (.txt/.md/.csv/.json/.docx)
  POST /rehydrate-doc  {filename, content_b64} -> {filename, content_b64}
  POST /redact-image   {filename, content_b64} -> {filename, content_b64}   (caviardage OCR, destructif)
  POST /redact-pdf     {filename, content_b64} -> {filename, content_b64}   (PDF -> pages caviardées, destructif)
  POST /detect         {text}                  -> {spans:[{start,end,type,text}]}
  GET  /health

Config (env) : PII_PROTECT_KEY, PII_NER_MODEL, PII_NER_THRESHOLD, PORT.
"""
import base64, io, json, os
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer

import crypto, analyzer, ner, image_redact

PORT = int(os.environ.get("PORT", "8077"))

# ╔══════════════════════════════════════════════════════════════════════════╗
# ║  CHOISIR CE QUI EST CHIFFRÉ (par TYPE de PII)                              ║
# ║  Par défaut : tout ce qui est détecté est chiffré (secure-by-default).     ║
# ║    PII_PROTECT_ENTITIES = "NOM_PERSONNE,PRENOM_PERSONNE,NUM_SECURITE..."    ║
# ║        -> allow-list : SEULS ces types sont chiffrés (le reste reste clair) ║
# ║    PII_PROTECT_SKIP = "VILLE,DATE,CODE_POSTAL"                              ║
# ║        -> deny-list : ces types restent en clair (tout le reste chiffré)    ║
# ║  (le filtrage par CHAMP se fait côté extension via PII_PROTECT_FIELDS)      ║
# ╚══════════════════════════════════════════════════════════════════════════╝
def _types_env(name):
    return {t.strip().upper() for t in os.environ.get(name, "").split(",") if t.strip()}
ONLY_TYPES = _types_env("PII_PROTECT_ENTITIES")   # vide = pas de restriction
SKIP_TYPES = _types_env("PII_PROTECT_SKIP")

def _should_encrypt(entity_type):
    if ONLY_TYPES and entity_type not in ONLY_TYPES:
        return False
    return entity_type not in SKIP_TYPES

def anonymize(text):
    spans = [sp for sp in analyzer.detect(text) if _should_encrypt(sp[2])]
    out, entities = text, []
    for s, e, t in sorted(spans, key=lambda x: x[0], reverse=True):
        out = out[:s] + crypto.marker(t, text[s:e]) + out[e:]
        entities.append(t)
    return out, list(reversed(entities))

# ---------------------------------------------------------------- documents texte
def _docx_transform(data, fn):
    from docx import Document
    d = Document(io.BytesIO(data))
    def paras(doc):
        for p in doc.paragraphs: yield p
        for t in doc.tables:
            for row in t.rows:
                for cell in row.cells:
                    for p in cell.paragraphs: yield p
    for p in paras(d):
        if p.text.strip():
            new = fn(p.text)
            if new != p.text:
                runs = list(p.runs)
                for r in runs: r.text = ""
                if runs: runs[0].text = new
                else: p.add_run(new)
    out = io.BytesIO(); d.save(out); return out.getvalue()

def process_doc(filename, data, mode):
    ext = (filename or "").lower().rsplit(".", 1)[-1]
    fn = (lambda t: anonymize(t)[0]) if mode == "anonymize" else (lambda t: crypto.rehydrate(t)[0])
    if ext == "docx":
        return _docx_transform(data, fn)
    return fn(data.decode("utf-8", "replace")).encode("utf-8")

# ---------------------------------------------------------------- routes POST
# Chaque handler reçoit le body (dict JSON) et renvoie un dict (sérialisé en réponse 200).
def _file_op(body, fn):
    """Transforme un fichier base64 -> base64 en conservant le nom."""
    out = fn(base64.b64decode(body.get("content_b64", "")))
    return {"filename": body.get("filename", ""), "content_b64": base64.b64encode(out).decode()}

def _anonymize(b):
    out, ents = anonymize(b.get("text", ""))
    return {"text": out, "count": len(ents), "entities": ents}

def _encrypt(b):
    # Chiffre la VALEUR ENTIÈRE en un marqueur (ex. nom de fichier : on protège tout, pas par PII).
    t = b.get("text", "")
    return {"text": crypto.marker("TEXT", t) if t else t, "count": 1 if t else 0}

def _rehydrate(b):
    out, c = crypto.rehydrate(b.get("text", ""))
    return {"text": out, "count": c}

def _detect(b):
    t = b.get("text", "")
    return {"spans": [{"start": s, "end": e, "type": ty, "text": t[s:e]} for s, e, ty in analyzer.detect(t)]}

POST_ROUTES = {
    "/anonymize":     _anonymize,
    "/encrypt":       _encrypt,
    "/rehydrate":     _rehydrate,
    "/detect":        _detect,
    "/anonymize-doc": lambda b: _file_op(b, lambda d: process_doc(b.get("filename", ""), d, "anonymize")),
    "/rehydrate-doc": lambda b: _file_op(b, lambda d: process_doc(b.get("filename", ""), d, "rehydrate")),
    "/redact-image":  lambda b: _file_op(b, image_redact.redact_image),
    "/redact-pdf":    lambda b: _file_op(b, image_redact.redact_pdf),
}

# ---------------------------------------------------------------- HTTP
class H(BaseHTTPRequestHandler):
    def _send(self, code, obj):
        b = json.dumps(obj, ensure_ascii=False).encode()
        self.send_response(code)
        self.send_header("Content-Type", "application/json; charset=utf-8")
        self.send_header("Content-Length", str(len(b)))
        self.end_headers(); self.wfile.write(b)
    def log_message(self, *a): pass
    def do_GET(self):
        if self.path == "/health":
            self._send(200, {"status": "ok", "model": ner.MODEL})
        else:
            self._send(404, {"error": "not found"})
    def do_POST(self):
        handler = POST_ROUTES.get(self.path)
        if handler is None:
            return self._send(404, {"error": "not found"})
        try:
            body = json.loads(self.rfile.read(int(self.headers.get("Content-Length", 0))) or b"{}")
            self._send(200, handler(body))
        except Exception as ex:
            self._send(500, {"error": str(ex)})

if __name__ == "__main__":
    print("Chargement de Presidio + du modèle NER…", flush=True)
    analyzer.warmup()
    print(f"PII engine prêt sur :{PORT}", flush=True)
    ThreadingHTTPServer(("0.0.0.0", PORT), H).serve_forever()
